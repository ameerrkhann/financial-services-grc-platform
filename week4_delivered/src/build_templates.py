# build_templates.py
# Builds the two Excel working templates: a risk register and a
# cross-framework control matrix.

import os
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

from grc_data import SCENARIOS

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "out")
TPL = os.path.join(OUT, "templates")
os.makedirs(TPL, exist_ok=True)

HDR_BG   = "14508C"
HDR_FG   = "FFFFFF"
BAND     = "EEF1F6"
CRITICAL = "F4C7C3"
HIGH     = "FBE2C7"
MEDIUM   = "FCF0C2"
LOW      = "D4EDDA"
INK      = "14171F"

THIN = Side(style="thin", color="C9D0DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header_row(ws, headers, row=1, height=34):
    ws.row_dimensions[row].height = height
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color=HDR_FG, size=10)
        c.fill = PatternFill("solid", fgColor=HDR_BG)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_body(ws, first_row, last_row, ncols, wrap_cols=()):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=9.5, color=INK)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(c in wrap_cols),
                horizontal="left",
            )
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=BAND)


def add_dv(ws, formula, cell_range, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showDropDown=False, showErrorMessage=True,
                        errorTitle="Invalid value",
                        error="Pick a value from the list.")
    if prompt:
        dv.promptTitle = "Guidance"
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


# ──────────────────────────────────────────────────────────────────────
# Risk register
# ──────────────────────────────────────────────────────────────────────

RR_HEADERS = [
    "Risk ID", "Risk Title", "Threat Source", "Asset / Process",
    "Likelihood (1-5)", "Impact (1-5)", "Inherent Risk", "Inherent Rating",
    "Existing Controls", "Control Effectiveness %", "Residual Risk",
    "Residual Rating", "Risk Owner", "Treatment", "Treatment Plan",
    "Target Date", "Status", "ALE (FAIR, $)", "NIST CSF 2.0 Ref", "OSFI Ref",
]
RR_WIDTHS = [9, 30, 18, 22, 12, 11, 11, 13, 34, 13, 11, 13, 16, 13, 38, 12, 13, 14, 16, 30]
BLANK_ROWS = 20

# recommended controls keyed by short scenario name
SCENARIO_CONTROLS = {v["short"]: v["controls"] for v in SCENARIOS.values()}


def build_risk_register(scen, gaps):
    wb = Workbook()

    # ── Instructions ──────────────────────────────────────────────────
    ins = wb.active
    ins.title = "Instructions"
    set_widths(ins, [26, 96])
    ins["A1"] = "Cyber Risk Register"
    ins["A1"].font = Font(bold=True, size=16, color=HDR_BG)
    ins["A2"] = "First National Bank (Fictional) — simulated data for portfolio demonstration"
    ins["A2"].font = Font(size=10, italic=True, color="4A5163")

    rows = [
        ("", ""),
        ("How to use", "Add one row per risk on the Risk Register sheet. Grey columns calculate "
                       "automatically — do not type over them."),
        ("", ""),
        ("Column", "What it means"),
        ("Risk ID", "Unique identifier, e.g. R-001. Keep it stable once issued."),
        ("Risk Title", "One line naming the risk event, not the control gap."),
        ("Threat Source", "Who or what causes it: external attacker, insider, vendor, system failure."),
        ("Asset / Process", "What is affected — the system, data set or business process."),
        ("Likelihood (1-5)", "1 Rare · 2 Unlikely · 3 Possible · 4 Likely · 5 Almost certain."),
        ("Impact (1-5)", "1 Negligible · 2 Minor · 3 Moderate · 4 Major · 5 Severe."),
        ("Inherent Risk", "CALCULATED: Likelihood × Impact, before controls."),
        ("Inherent Rating", "CALCULATED: 15+ Critical · 10-14 High · 5-9 Medium · below 5 Low."),
        ("Existing Controls", "Controls already operating against this risk."),
        ("Control Effectiveness %", "How much of the risk those controls remove. Evidence-based where "
                                    "possible; state the basis in Treatment Plan if assumed."),
        ("Residual Risk", "CALCULATED: Inherent Risk × (1 − Control Effectiveness)."),
        ("Residual Rating", "CALCULATED on the same bands as Inherent Rating."),
        ("Risk Owner", "The named accountable person. Not a team."),
        ("Treatment", "Accept · Mitigate · Transfer · Avoid."),
        ("Treatment Plan", "What will be done, by whom. Specific enough to track."),
        ("Target Date", "When the treatment completes."),
        ("Status", "Open · In Progress · Monitoring · Closed."),
        ("ALE (FAIR, $)", "Annualised loss expectancy from the quantitative risk engine, where the "
                          "risk has been modelled."),
        ("NIST CSF 2.0 Ref", "Relevant CSF 2.0 subcategory, e.g. GV.PO-01."),
        ("OSFI Ref", "Relevant OSFI guideline: B-10, B-13 or E-21."),
        ("", ""),
        ("Rating bands", "Critical 15-25 · High 10-14 · Medium 5-9 · Low 1-4"),
        ("Note on the examples", "The five pre-filled rows are the modelled FAIR scenarios. Their ALE "
                                 "figures come from the quantitative risk engine. Replace or extend them."),
    ]
    for i, (a, b) in enumerate(rows, start=4):
        ins[f"A{i}"] = a
        ins[f"B{i}"] = b
        ins[f"A{i}"].font = Font(bold=True, size=9.5, color=INK)
        ins[f"B{i}"].font = Font(size=9.5, color=INK)
        ins[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ins.row_dimensions[i].height = 26 if len(str(b)) > 90 else 15
    ins["A7"].fill = PatternFill("solid", fgColor=HDR_BG)
    ins["B7"].fill = PatternFill("solid", fgColor=HDR_BG)
    ins["A7"].font = Font(bold=True, color=HDR_FG, size=10)
    ins["B7"].font = Font(bold=True, color=HDR_FG, size=10)

    # ── Register ──────────────────────────────────────────────────────
    ws = wb.create_sheet("Risk Register")
    header_row(ws, RR_HEADERS)
    set_widths(ws, RR_WIDTHS)

    likelihood_map = {"Ransomware": 4, "Data Breach": 4, "Cloud Misconfig": 5,
                      "Vendor Failure": 3, "Insider Threat": 3}
    impact_map = {"Ransomware": 5, "Data Breach": 5, "Cloud Misconfig": 4,
                  "Vendor Failure": 4, "Insider Threat": 4}

    r = 2
    for s in scen.sort_values("ale", ascending=False).itertuples():
        short = s.scenario_short
        eff = int(round(s.control_effectiveness * 100))
        ws.cell(row=r, column=1, value=f"R-{s.id:03d}" if hasattr(s, "id") else f"R-{r-1:03d}")
        ws.cell(row=r, column=2, value=s.scenario_name)
        ws.cell(row=r, column=3, value={"Ransomware": "External attacker",
                                        "Data Breach": "External attacker",
                                        "Insider Threat": "Malicious or negligent insider",
                                        "Vendor Failure": "Third-party provider",
                                        "Cloud Misconfig": "Internal misconfiguration"}[short])
        ws.cell(row=r, column=4, value={"Ransomware": "Core banking platform",
                                        "Data Breach": "Customer PII datastore",
                                        "Insider Threat": "Privileged access to customer records",
                                        "Vendor Failure": "Critical vendor services",
                                        "Cloud Misconfig": "Cloud storage and IAM"}[short])
        ws.cell(row=r, column=5, value=likelihood_map[short])
        ws.cell(row=r, column=6, value=impact_map[short])
        ws.cell(row=r, column=7, value=f"=E{r}*F{r}")
        ws.cell(row=r, column=8, value=f'=IF(G{r}="","",IF(G{r}>=15,"Critical",IF(G{r}>=10,"High",IF(G{r}>=5,"Medium","Low"))))')
        ws.cell(row=r, column=9, value="; ".join(SCENARIO_CONTROLS[short]))
        ws.cell(row=r, column=10, value=eff / 100)
        ws.cell(row=r, column=11, value=f"=ROUND(G{r}*(1-J{r}),1)")
        ws.cell(row=r, column=12, value=f'=IF(K{r}="","",IF(K{r}>=15,"Critical",IF(K{r}>=10,"High",IF(K{r}>=5,"Medium","Low"))))')
        ws.cell(row=r, column=13, value="CISO")
        ws.cell(row=r, column=14, value="Mitigate")
        ws.cell(row=r, column=15, value=f"Implement recommended controls "
                                        f"(annual cost ${s.control_cost:,.0f}).")
        ws.cell(row=r, column=16, value="2027-03-31")
        ws.cell(row=r, column=17, value="Open")
        ws.cell(row=r, column=18, value=float(s.ale))
        ws.cell(row=r, column=19, value=s.csf_ref)
        ws.cell(row=r, column=20, value=s.osfi_ref)
        r += 1

    first_blank = r
    last = r + BLANK_ROWS - 1
    for rr in range(first_blank, last + 1):
        ws.cell(row=rr, column=7, value=f"=IF(COUNT(E{rr}:F{rr})=2,E{rr}*F{rr},\"\")")
        ws.cell(row=rr, column=8, value=f'=IF(G{rr}="","",IF(G{rr}>=15,"Critical",IF(G{rr}>=10,"High",IF(G{rr}>=5,"Medium","Low"))))')
        ws.cell(row=rr, column=11, value=f'=IF(OR(G{rr}="",J{rr}=""),"",ROUND(G{rr}*(1-J{rr}),1))')
        ws.cell(row=rr, column=12, value=f'=IF(K{rr}="","",IF(K{rr}>=15,"Critical",IF(K{rr}>=10,"High",IF(K{rr}>=5,"Medium","Low"))))')

    style_body(ws, 2, last, len(RR_HEADERS), wrap_cols={2, 9, 15, 20})

    # formats
    for rr in range(2, last + 1):
        ws.cell(row=rr, column=10).number_format = "0%"
        ws.cell(row=rr, column=18).number_format = '$#,##0'
        for c in (5, 6, 7, 11):
            ws.cell(row=rr, column=c).alignment = Alignment(horizontal="center", vertical="top")
        # calculated columns visually distinct
        for c in (7, 8, 11, 12):
            ws.cell(row=rr, column=c).font = Font(size=9.5, italic=True, color="3A4254")

    # validation
    add_dv(ws, '"1,2,3,4,5"', f"E2:E{last}", "1 Rare · 5 Almost certain")
    add_dv(ws, '"1,2,3,4,5"', f"F2:F{last}", "1 Negligible · 5 Severe")
    add_dv(ws, '"Accept,Mitigate,Transfer,Avoid"', f"N2:N{last}")
    add_dv(ws, '"Open,In Progress,Monitoring,Closed"', f"Q2:Q{last}")
    add_dv(ws, '"B-10,B-13,E-21"', f"T{first_blank}:T{last}")

    # conditional formatting on both rating columns
    for colrange in (f"H2:H{last}", f"L2:L{last}"):
        for word, fill in (("Critical", CRITICAL), ("High", HIGH),
                           ("Medium", MEDIUM), ("Low", LOW)):
            ws.conditional_formatting.add(colrange, CellIsRule(
                operator="equal", formula=[f'"{word}"'],
                fill=PatternFill("solid", fgColor=fill)))

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(RR_HEADERS))}{last}"

    path = os.path.join(TPL, "risk_register.xlsx")
    wb.save(path)
    print(f"  wrote {os.path.relpath(path, OUT)}  ({r-2} examples + {BLANK_ROWS} blank rows)")
    return path


# ──────────────────────────────────────────────────────────────────────
# Control matrix
# ──────────────────────────────────────────────────────────────────────

CM_HEADERS = [
    "Control ID", "Control Description", "CSF Function", "NIST CSF 2.0",
    "ISO 27001:2022", "SOC 2 TSC", "OSFI Guideline", "Control Owner",
    "Test Frequency", "Last Tested", "Test Result", "Evidence Link",
]
CM_WIDTHS = [11, 52, 13, 14, 16, 11, 34, 16, 14, 13, 19, 24]


def build_control_matrix(mapping):
    wb = Workbook()
    ws = wb.active
    ws.title = "Control Matrix"
    header_row(ws, CM_HEADERS)
    set_widths(ws, CM_WIDTHS)

    r = 2
    for i, m in enumerate(mapping.itertuples(), start=1):
        ws.cell(row=r, column=1, value=f"C-{i:03d}")
        ws.cell(row=r, column=2, value=m.control)
        ws.cell(row=r, column=3, value=m.function)
        ws.cell(row=r, column=4, value=m.csf_ref)
        ws.cell(row=r, column=5, value=m.iso_ref)
        ws.cell(row=r, column=6, value=m.soc2_ref)
        ws.cell(row=r, column=7, value=m.osfi_ref)
        ws.cell(row=r, column=8, value="")
        ws.cell(row=r, column=9, value=m.frequency)
        ws.cell(row=r, column=10, value="")
        ws.cell(row=r, column=11, value="Not Tested")
        ws.cell(row=r, column=12, value="")
        r += 1
    last = r - 1

    style_body(ws, 2, last, len(CM_HEADERS), wrap_cols={2, 7})
    add_dv(ws, '"Effective,Partially Effective,Ineffective,Not Tested"', f"K2:K{last}")
    add_dv(ws, '"Continuous,Monthly,Quarterly,Annual,Per incident"', f"I2:I{last}")

    for word, fill in (("Effective", LOW), ("Partially Effective", MEDIUM),
                       ("Ineffective", CRITICAL), ("Not Tested", "E4E7EE")):
        ws.conditional_formatting.add(f"K2:K{last}", CellIsRule(
            operator="equal", formula=[f'"{word}"'],
            fill=PatternFill("solid", fgColor=fill)))

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CM_HEADERS))}{last}"

    # ── Summary ───────────────────────────────────────────────────────
    sm = wb.create_sheet("Summary")
    set_widths(sm, [26, 14, 14, 14, 14, 14])
    sm["A1"] = "Control Coverage Summary"
    sm["A1"].font = Font(bold=True, size=14, color=HDR_BG)
    sm["A2"] = "Counts update automatically from the Control Matrix sheet."
    sm["A2"].font = Font(size=9.5, italic=True, color="4A5163")

    header_row(sm, ["CSF Function", "Controls", "Effective", "Partially Effective",
                    "Ineffective", "Not Tested"], row=4, height=30)
    funcs = ["Govern", "Identify", "Protect", "Detect", "Respond", "Recover"]
    row = 5
    for f in funcs:
        sm.cell(row=row, column=1, value=f)
        sm.cell(row=row, column=2,
                value=f"=COUNTIF('Control Matrix'!$C$2:$C${last},$A{row})")
        for j, res in enumerate(["Effective", "Partially Effective",
                                 "Ineffective", "Not Tested"], start=3):
            sm.cell(row=row, column=j,
                    value=f"=COUNTIFS('Control Matrix'!$C$2:$C${last},$A{row},"
                          f"'Control Matrix'!$K$2:$K${last},\"{res}\")")
        row += 1
    sm.cell(row=row, column=1, value="Total")
    for j in range(2, 7):
        c = get_column_letter(j)
        sm.cell(row=row, column=j, value=f"=SUM({c}5:{c}{row-1})")
    style_body(sm, 5, row, 6)
    for rr in range(5, row + 1):
        sm.cell(row=rr, column=1).font = Font(bold=True, size=9.5, color=INK)
        for c in range(2, 7):
            sm.cell(row=rr, column=c).alignment = Alignment(horizontal="center")

    path = os.path.join(TPL, "control_matrix.xlsx")
    wb.save(path)
    print(f"  wrote {os.path.relpath(path, OUT)}  ({last-1} controls)")
    return path


if __name__ == "__main__":
    scen = pd.read_json(os.path.join(OUT, "RiskScenarios.json"))
    gaps = pd.read_json(os.path.join(OUT, "ControlGaps.json"))
    mapping = pd.read_json(os.path.join(OUT, "ControlMapping.json"))
    print("Building Excel templates...")
    build_risk_register(scen, gaps)
    build_control_matrix(mapping)
