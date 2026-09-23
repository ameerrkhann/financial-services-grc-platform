# build_dataset.py
# Builds the complete, reproducible GRC demo dataset and writes the
# Power BI export workbook.

import os
import json
from datetime import date, timedelta

import numpy as np
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from fair_engine import run_fair_simulation, compute_rosi
from grc_data import (
    ORG_NAME, ASSESSOR, SEED, N_SIMULATIONS,
    CSF_FUNCTIONS, CONTROL_GAPS, CONTROL_MAPPING, SCENARIOS,
    RISK_TIERS, DEMO_VENDORS, VENDOR_CATEGORIES,
)

TODAY = date(2026, 9, 21)
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
os.makedirs(OUT_DIR, exist_ok=True)

# ──────────────────────────────────────────────────────────────────────
# Vendor questionnaire — weights mirror the Streamlit app
# ──────────────────────────────────────────────────────────────────────

UNIVERSAL_WEIGHTS = [
    ("U01", "Governance", 3, "Documented information security policy approved by senior leadership"),
    ("U02", "Governance", 3, "Recognised security certification maintained (ISO 27001 / SOC 2 Type II)"),
    ("U03", "Incident Response", 3, "Incident response plan includes client notification within 24 hours"),
    ("U04", "Incident Response", 2, "Security incident history disclosed with post-incident reporting"),
    ("U05", "Access Control", 3, "MFA enforced on all systems with access to client data"),
    ("U06", "Access Control", 2, "Least-privilege access applied and reviewed"),
    ("U07", "Data Protection", 3, "Client data encrypted at rest (AES-256) and in transit (TLS 1.2+)"),
    ("U08", "Data Protection", 2, "Documented data retention and secure deletion policy"),
    ("U09", "Business Continuity", 3, "Tested business continuity plan with defined recovery time objectives"),
    ("U10", "Business Continuity", 2, "Backups performed regularly and restore-tested"),
    ("U11", "Vulnerability Management", 2, "Independent penetration testing conducted at least annually"),
    ("U12", "Vulnerability Management", 2, "Formal vulnerability management process with patching SLAs"),
    ("U13", "Subcontracting", 2, "Fourth parties subject to equivalent security requirements"),
    ("U14", "Audit Rights", 2, "Client permitted to audit or review independent assurance reports"),
]

SERVICE_WEIGHTS = {
    "Cloud Provider": [
        ("CL01", "Cloud Security", 3, "Data residency guarantees for Canadian customer data"),
        ("CL02", "Cloud Security", 2, "Cloud security posture management with continuous misconfiguration scanning"),
        ("CL03", "Cloud Security", 3, "Documented exit strategy enabling data migration off-platform"),
        ("CL04", "Cloud Security", 3, "Logical tenant isolation between client environments"),
    ],
    "Payment Processor": [
        ("PP01", "Payment Security", 3, "PCI DSS Level 1 compliant with current Report on Compliance"),
        ("PP02", "Payment Security", 3, "Tokenisation or encryption protecting card data in processing"),
        ("PP03", "Payment Security", 2, "Real-time fraud detection and transaction monitoring"),
        ("PP04", "Payment Security", 3, "Contractual availability SLA for transaction processing"),
    ],
    "Software / SaaS": [
        ("SW01", "Software Security", 2, "Secure software development lifecycle with stage-gate security testing"),
        ("SW02", "Software Security", 2, "SAST and DAST performed on all releases"),
        ("SW03", "Software Security", 2, "Defined SLA for critical vulnerability remediation"),
        ("SW04", "Software Security", 2, "Software bill of materials provided for third-party components"),
    ],
    "Data Analytics / AI": [
        ("DA01", "Data Governance", 3, "Client data not used for model training without documented consent"),
        ("DA02", "Data Governance", 2, "Documented data minimisation policy"),
        ("DA03", "Data Governance", 2, "Model decisions explainable and auditable by the client"),
        ("DA04", "Data Governance", 2, "Bias and fairness testing on models informing customer decisions"),
    ],
    "IT Infrastructure": [
        ("IT01", "Network Security", 2, "Network segmentation isolating client environments"),
        ("IT02", "Network Security", 3, "Remote access via encrypted VPN or zero-trust network access"),
        ("IT03", "Physical Security", 2, "Data centres certified to SOC 2 Type II or ISO 27001"),
        ("IT04", "Network Security", 2, "24/7 security operations centre with defined escalation"),
    ],
}


def build_vendor_answers(service_type, target_score, rng):
    """
    Produces an answer set whose weighted score lands close to target_score.
    Higher-weight questions are answered favourably first, which mirrors how a
    mature vendor tends to have the fundamental controls in place.
    """
    questions = UNIVERSAL_WEIGHTS + SERVICE_WEIGHTS[service_type]
    max_points = sum(w * 2 for _, _, w, _ in questions)
    target_points = round(max_points * target_score / 100.0)

    # Vendors tend to have the high-weight fundamentals covered, but coverage is
    # never perfectly ordered by importance. Rank by weight plus noise so some
    # high-weight controls genuinely fail — that is what produces critical gaps.
    order = sorted(range(len(questions)),
                   key=lambda i: -(questions[i][2] + rng.normal(0, 1.1)))

    answers = {}
    earned = 0
    for i in order:
        qid, cat, weight, text = questions[i]
        remaining = target_points - earned
        if remaining >= weight * 2:
            ans, pts = "Yes", weight * 2
        elif remaining >= weight:
            ans, pts = "Partial", weight
        else:
            ans, pts = "No", 0
        earned += pts
        answers[qid] = {"answer": ans, "weight": weight,
                        "category": cat, "question": text,
                        "points_earned": pts, "points_possible": weight * 2}

    score = int(round(earned / max_points * 100))
    return answers, score


def tier_for_score(score):
    for name, info in RISK_TIERS.items():
        if info["min"] <= score <= info["max"]:
            return name, info
    return "Critical", RISK_TIERS["Critical"]


def build_vendors():
    rng = np.random.default_rng(SEED)
    vendors, responses = [], []
    for v in DEMO_VENDORS:
        answers, score = build_vendor_answers(v["service"], v["target"], rng)
        tier, info = tier_for_score(score)
        assess_date = TODAY - timedelta(days=v["days_ago"])
        due = assess_date + timedelta(days=info["reassess_days"])
        overdue = due < TODAY

        gaps = [a for a in answers.values() if a["answer"] in ("No", "Partial")]
        crit = [a for a in gaps if a["answer"] == "No" and a["weight"] == 3]

        vendors.append({
            "vendor_name": v["name"],
            "service_type": v["service"],
            "criticality": v["criticality"],
            "risk_tier": tier,
            "score": score,
            "gap_count": len(gaps),
            "critical_gaps": len(crit),
            "assessor": ASSESSOR,
            "assessment_date": assess_date.isoformat(),
            "days_since_assessment": (TODAY - assess_date).days,
            "reassessment_due": due.isoformat(),
            "is_overdue": "TRUE" if overdue else "FALSE",
            "is_overdue_flag": 1 if overdue else 0,   # numeric so a Power BI SUM card works
            "is_critical_flag": 1 if tier == "Critical" else 0,
            "recommended_action": info["action"],
        })
        for qid, a in answers.items():
            responses.append({
                "vendor_name": v["name"], "question_id": qid,
                "category": a["category"], "question_text": a["question"],
                "answer": a["answer"], "weight": a["weight"],
                "points_earned": a["points_earned"],
                "points_possible": a["points_possible"],
            })
    return pd.DataFrame(vendors), pd.DataFrame(responses)


def band_likelihood(lam):
    return "Low" if lam < 1.0 else ("Medium" if lam <= 2.0 else "High")


def band_impact(loss_high):
    return "Medium" if loss_high < 5_000_000 else ("High" if loss_high < 9_000_000 else "Critical")


def build_scenarios():
    rows, lec_rows, curves = [], [], {}
    for key, s in SCENARIOS.items():
        r = run_fair_simulation(s["loss_low"], s["loss_high"],
                                s["freq_low"], s["freq_high"],
                                n=N_SIMULATIONS, seed=SEED + s["id"])
        roi = compute_rosi(r["ale"], s["control_cost"], s["control_effectiveness"])
        lam = (s["freq_low"] + s["freq_high"]) / 2

        rows.append({
            "scenario_key": key,
            "scenario_name": s["name"],
            "scenario_short": s["short"],
            "loss_low": s["loss_low"], "loss_high": s["loss_high"],
            "freq_low": s["freq_low"], "freq_high": s["freq_high"],
            "avg_frequency": lam,
            "ale": round(r["ale"]),
            "median": round(r["median"]),
            "percentile_90": round(r["percentile_90"]),
            "percentile_95": round(r["percentile_95"]),
            "prob_over_1m": round(r["prob_over_1m"], 2),
            "prob_over_5m": round(r["prob_over_5m"], 2),
            "control_cost": s["control_cost"],
            "control_effectiveness": s["control_effectiveness"],
            "residual_ale": round(roi["residual_ale"]),
            "risk_reduction": round(roi["risk_reduction"]),
            "rosi_pct": round(roi["rosi_pct"], 1),
            "likelihood_band": band_likelihood(lam),
            "impact_band": band_impact(s["loss_high"]),
            "osfi_ref": s["osfi_ref"],
            "csf_ref": s["csf_ref"],
            "date_run": TODAY.isoformat(),
        })

        losses = r["annual_losses"]
        curves[key] = (s, r)
        max_loss = np.percentile(losses, 99.5)
        for lp in np.linspace(0, max_loss, 100):
            lec_rows.append({
                "scenario_name": s["name"],
                "loss_value": round(float(lp)),
                "exceedance_pct": round(float(np.mean(losses > lp) * 100), 2),
            })
    return pd.DataFrame(rows), pd.DataFrame(lec_rows), curves


def build_csf():
    scores, gaps = [], []
    for name, f in sorted(CSF_FUNCTIONS.items(), key=lambda kv: kv[1]["order"]):
        scores.append({
            "function_order": f["order"],
            "function_name": name,
            "function_code": f["code"],
            "score": f["score"],
            "target_score": f["target"],
            "gap_to_target": f["target"] - f["score"],
            "status": "OK" if f["score"] >= 3 else "GAP",
            "rationale": f["rationale"],
        })
    for name, g in sorted(CONTROL_GAPS.items(), key=lambda kv: kv[1]["risk_rank"]):
        gaps.append({
            # 'note' is retained so an existing Power BI query that was built when
            # this sheet was a placeholder still refreshes. Power Query fails on a
            # REMOVED column but tolerates added ones.
            "note": f"{name} gap — {g['priority']} priority",
            "risk_rank": g["risk_rank"],
            "function_name": name,
            "gap_description": g["gap"],
            "priority": g["priority"],
            "csf_ref": g["csf_ref"],
            "iso27001_ref": g["iso_ref"],
            "soc2_ref": g["soc2_ref"],
            "osfi_ref": g["osfi_ref"],
            "effort": g["effort"],
            "effort_weeks": g["effort_weeks"],
            "business_impact": g["impact"],
            "quick_win": g["quick_win"],
            "remediation": g["remediation"],
        })
    return pd.DataFrame(scores), pd.DataFrame(gaps)


def build_kpis(scen_df, csf_df, vendor_df):
    """
    compliance_posture_score blends the two maturity signals into one 0-100
    figure: 50% average vendor questionnaire score, 50% NIST CSF overall
    maturity expressed as a percentage of the 5-point scale.
    """
    avg_vendor = float(vendor_df["score"].mean())
    csf_overall = float(csf_df["score"].mean())
    posture = round(0.5 * avg_vendor + 0.5 * (csf_overall / 5 * 100), 1)
    return pd.DataFrame([{
        "org_name": ORG_NAME,
        "report_date": TODAY.isoformat(),
        "vendors_assessed": int(len(vendor_df)),
        "critical_vendors": int((vendor_df["risk_tier"] == "Critical").sum()),
        "high_risk_vendors": int((vendor_df["risk_tier"].isin(["Critical", "High"])).sum()),
        "overdue_vendors": int((vendor_df["is_overdue"] == "TRUE").sum()),
        "avg_vendor_score": round(avg_vendor, 1),
        "csf_overall_score": round(csf_overall, 2),
        "csf_target_score": float(csf_df["target_score"].mean()),
        "open_control_gaps": 0,  # filled by caller
        "scenarios_assessed": int(len(scen_df)),
        "total_ale": int(scen_df["ale"].sum()),
        "total_residual_ale": int(scen_df["residual_ale"].sum()),
        "total_control_cost": int(scen_df["control_cost"].sum()),
        "total_risk_reduction": int(scen_df["risk_reduction"].sum()),
        "compliance_posture_score": posture,
    }])


def write_workbook(sheets, path):
    """Writes each DataFrame to its own sheet and registers it as an Excel Table."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            if df.empty:
                df = pd.DataFrame({"note": ["No data"]})
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
            tbl = Table(displayName=f"{name}_tbl", ref=ref)
            tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tbl)
            for i, col in enumerate(df.columns, start=1):
                widths = [len(str(col))] + [len(str(v)) for v in df[col].head(30)]
                ws.column_dimensions[get_column_letter(i)].width = min(max(max(widths) + 2, 10), 45)


def main():
    print(f"Building GRC demo dataset for {ORG_NAME}")
    print(f"  seed={SEED}  simulations={N_SIMULATIONS:,}  as-of {TODAY.isoformat()}\n")

    scen_df, lec_df, curves = build_scenarios()
    csf_df, gaps_df = build_csf()
    vendor_df, resp_df = build_vendors()

    kpi_df = build_kpis(scen_df, csf_df, vendor_df)
    kpi_df.loc[0, "open_control_gaps"] = int(len(gaps_df))

    sheets = {
        "RiskScenarios": scen_df,
        "LossExceedance": lec_df,
        "CSFScores": csf_df,
        "ControlGaps": gaps_df,
        "VendorAssessments": vendor_df,
        "VendorResponses": resp_df,
        "VendorTierSummary": vendor_df.groupby("risk_tier").size()
            .reindex(["Critical", "High", "Medium", "Low"], fill_value=0)
            .reset_index(name="vendor_count"),
        "ControlMapping": pd.DataFrame(CONTROL_MAPPING),
        "KPISummary": kpi_df,
    }

    xlsx = os.path.join(OUT_DIR, "grc_powerbi_data.xlsx")
    write_workbook(sheets, xlsx)

    for name, df in sheets.items():
        print(f"  {name:<20} {len(df):>5} rows")

    # Persist the raw simulation for chart and PDF generation
    np.savez_compressed(
        os.path.join(OUT_DIR, "simulations.npz"),
        **{k: v[1]["annual_losses"] for k, v in curves.items()}
    )
    for name, df in sheets.items():
        df.to_json(os.path.join(OUT_DIR, f"{name}.json"), orient="records")

    k = kpi_df.iloc[0]
    print(f"\n  Total ALE            ${k['total_ale']:,}")
    print(f"  Residual ALE         ${k['total_residual_ale']:,}")
    print(f"  Control investment   ${k['total_control_cost']:,}")
    print(f"  CSF overall          {k['csf_overall_score']}/5")
    print(f"  Vendors              {k['vendors_assessed']} "
          f"({k['critical_vendors']} critical, {k['overdue_vendors']} overdue)")
    print(f"  Posture score        {k['compliance_posture_score']}/100")
    print(f"\n  Workbook: {xlsx}")


if __name__ == "__main__":
    main()
