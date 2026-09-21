# src/reporting/export_for_powerbi.py
# Exports all GRC data from SQLite into one Excel file for Power BI.
# One file → multiple sheets, one per dashboard data source.
#
# The live Power BI report binds to the sheet names and to the "<Sheet>_tbl"
# Excel table names below. Adding columns is safe. Renaming or removing a sheet
# or a table breaks the published report, so don't.

import sys
import os
from datetime import date

import numpy as np
import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "../.."))

from src.database.db_manager import (
    get_connection,
    initialise_database,
    get_reassessment_days,
)
from src.risk_quantification.fair_engine import run_fair_simulation, seed_for
from src.risk_quantification.scenarios import SCENARIOS
from src.compliance.csf_data import DEFAULT_TARGET_SCORE

OUTPUT_DIR  = os.path.join(os.path.dirname(__file__), "../../dashboards")
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "grc_powerbi_data.xlsx")


def band_likelihood(avg_freq):
    if avg_freq < 1.0:
        return "Low"
    elif avg_freq <= 2.0:
        return "Medium"
    return "High"


def band_impact(loss_high):
    if loss_high < 5_000_000:
        return "Medium"
    elif loss_high < 9_000_000:
        return "High"
    return "Critical"


def get_risk_scenarios_df(conn):
    rows = conn.execute("""
        SELECT scenario_name, loss_low, loss_high, freq_low, freq_high,
               ale, median, percentile_90, percentile_95,
               prob_over_1m, prob_over_5m, control_cost,
               control_effectiveness, residual_ale, rosi
        FROM risk_scenarios
        WHERE id IN (SELECT MAX(id) FROM risk_scenarios GROUP BY scenario_key)
    """).fetchall()
    df = pd.DataFrame([dict(r) for r in rows])
    if not df.empty:
        df["avg_frequency"]   = (df["freq_low"] + df["freq_high"]) / 2
        df["likelihood_band"] = df["avg_frequency"].apply(band_likelihood)
        df["impact_band"]     = df["loss_high"].apply(band_impact)
        for col in ["ale", "median", "percentile_90", "percentile_95",
                    "residual_ale"]:
            df[col] = df[col].round(0)
        df["rosi_pct"] = (df["rosi"] * 100).round(0)
    return df


def get_lec_df():
    """
    Recomputes loss exceedance curve points for each scenario.

    Uses the same per-scenario seeds as run_scenarios.py so the curve in Power
    BI is drawn from the identical simulation that produced the stored ALE.
    """
    records = []
    for key, scenario in SCENARIOS.items():
        result   = run_fair_simulation(
            loss_low  = scenario["loss_low"],
            loss_high = scenario["loss_high"],
            freq_low  = scenario["freq_low"],
            freq_high = scenario["freq_high"],
            seed      = seed_for(scenario["id"]),
        )
        losses   = result["annual_losses"]
        max_loss = np.percentile(losses, 99.5)
        for lp in np.linspace(0, max_loss, 100):
            records.append({
                "scenario_name":  scenario["name"],
                "loss_value":     round(float(lp), 0),
                "exceedance_pct": round(float(np.mean(losses > lp) * 100), 2),
            })
    return pd.DataFrame(records)


def get_table_df(conn, query):
    rows = conn.execute(query).fetchall()
    return pd.DataFrame([dict(r) for r in rows])


def get_csf_scores_df(conn):
    """Latest CSF assessment: current score, target, and the gap between them."""
    df = get_table_df(conn, """
        SELECT function_name, score, target_score, rationale
        FROM function_scores
        WHERE assessment_id = (SELECT MAX(id) FROM assessments)
    """)
    if not df.empty:
        df["target_score"]  = df["target_score"].fillna(DEFAULT_TARGET_SCORE)
        df["gap_to_target"] = (df["target_score"] - df["score"]).clip(lower=0)
    return df


def get_vendor_assessments_df(conn):
    """
    Latest assessment per vendor, plus the reassessment ageing columns.

    reassessment_due is assessment_date + the window for that vendor's tier
    (RISK_TIERS[...]["reassessment_days"]), so the date shown in Power BI is
    derived from the same rule get_overdue_vendors() enforces.
    """
    df = get_table_df(conn, """
        SELECT vendor_name, service_type, criticality, risk_tier,
               score, gap_count, critical_gaps, assessor, assessment_date
        FROM vendor_assessments
        WHERE id IN (SELECT MAX(id) FROM vendor_assessments GROUP BY vendor_name)
    """)
    if df.empty:
        return df

    windows   = get_reassessment_days()
    assessed  = pd.to_datetime(df["assessment_date"])
    today     = pd.Timestamp(date.today())
    window    = df["risk_tier"].map(windows).fillna(365).astype(int)

    df["reassessment_window_days"] = window
    df["days_since_assessment"]    = (today - assessed).dt.days
    df["reassessment_due"]         = (
        (assessed + pd.to_timedelta(window, unit="D")).dt.strftime("%Y-%m-%d")
    )
    df["is_overdue"] = df["days_since_assessment"] > window
    return df


def get_kpi_summary_df(risk_df, csf_df, vendor_df):
    """
    One row of headline numbers for the dashboard's KPI cards.

    compliance_posture_score is a blended posture indicator on a 0-100 scale:

        compliance_posture_score = 0.5 * avg_vendor_score
                                 + 0.5 * (csf_overall_score / 5 * 100)

    Half the weight is the average vendor questionnaire score (already 0-100)
    and half is the NIST CSF overall maturity rescaled from /5 to /100. It is a
    presentation metric for trending one number over time, not a regulatory
    measure — no framework defines it.
    """
    avg_vendor_score = float(vendor_df["score"].mean()) if not vendor_df.empty else 0.0
    csf_overall      = float(csf_df["score"].mean())    if not csf_df.empty    else 0.0

    record = {
        "vendors_assessed":  int(len(vendor_df)),
        "critical_vendors":  int((vendor_df["risk_tier"] == "Critical").sum())
                             if not vendor_df.empty else 0,
        "overdue_vendors":   int(vendor_df["is_overdue"].sum())
                             if not vendor_df.empty else 0,
        "avg_vendor_score":  round(avg_vendor_score, 1),
        "csf_overall_score": round(csf_overall, 2),
        "total_ale":         round(float(risk_df["ale"].sum()), 0)
                             if not risk_df.empty else 0.0,
        "total_residual_ale": round(float(risk_df["residual_ale"].sum()), 0)
                             if not risk_df.empty else 0.0,
        "compliance_posture_score": round(
            0.5 * avg_vendor_score + 0.5 * (csf_overall / 5 * 100), 1
        ),
    }
    return pd.DataFrame([record])


def main():
    initialise_database()
    conn = get_connection()
    print("Exporting data for Power BI...\n")

    sheets = {}

    sheets["RiskScenarios"] = get_risk_scenarios_df(conn)
    print(f"  RiskScenarios     : {len(sheets['RiskScenarios'])} rows")

    sheets["LossExceedance"] = get_lec_df()
    print(f"  LossExceedance    : {len(sheets['LossExceedance'])} rows")

    sheets["CSFScores"] = get_csf_scores_df(conn)
    print(f"  CSFScores         : {len(sheets['CSFScores'])} rows")

    sheets["ControlGaps"] = get_table_df(conn, """
        SELECT function_name, gap_description, priority,
               nist_ref, iso27001_ref, soc2_ref
        FROM control_gaps
        WHERE assessment_id = (SELECT MAX(id) FROM assessments)
    """)
    print(f"  ControlGaps       : {len(sheets['ControlGaps'])} rows")

    sheets["VendorAssessments"] = get_vendor_assessments_df(conn)
    print(f"  VendorAssessments : {len(sheets['VendorAssessments'])} rows")

    sheets["VendorTierSummary"] = get_table_df(conn, """
        SELECT risk_tier, COUNT(*) as vendor_count
        FROM vendor_assessments
        WHERE id IN (SELECT MAX(id) FROM vendor_assessments GROUP BY vendor_name)
        GROUP BY risk_tier
    """)
    print(f"  VendorTierSummary : {len(sheets['VendorTierSummary'])} rows")

    sheets["KPISummary"] = get_kpi_summary_df(
        sheets["RiskScenarios"], sheets["CSFScores"], sheets["VendorAssessments"]
    )
    print(f"  KPISummary        : {len(sheets['KPISummary'])} rows")

    conn.close()

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for name, df in sheets.items():
            if df.empty:
                df = pd.DataFrame({"note": ["No data yet - run the relevant module first"]})
            df.to_excel(writer, sheet_name=name, index=False)

            # Format each sheet as a real Excel Table so Power BI can read it
            ws       = writer.sheets[name]
            n_rows   = len(df) + 1                       # data rows + header row
            last_col = get_column_letter(len(df.columns))
            ref      = f"A1:{last_col}{n_rows}"
            table    = Table(displayName=f"{name}_tbl", ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            ws.add_table(table)

    print(f"\n✅ Exported to: dashboards/grc_powerbi_data.xlsx")
    print("   Upload this file to Power BI at app.powerbi.com")


if __name__ == "__main__":
    main()